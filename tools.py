# -*- coding: utf-8 -*-
import os
import re
import sys
import json
import shutil
import time
import uuid
import fnmatch
import inspect
import ipaddress
import socket
import subprocess
import importlib.util
from html import unescape
from pathlib import Path
from urllib.parse import urlparse, urljoin, parse_qs

import requests
from backup_guard import BackupGuard, scrubbed_env

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CUSTOM_DIR = os.path.join(BASE_DIR, "custom_tools")
WORKSPACE_DIR = os.path.realpath(os.getenv("NODUS_WORKSPACE") or os.path.join(BASE_DIR, "workspace"))
os.makedirs(CUSTOM_DIR, exist_ok=True)
os.makedirs(WORKSPACE_DIR, exist_ok=True)

ERR = "[HATA] "          # Tum arac hatalari bu onekle baslar (basari kontrolu icin)
MAX_OUT = 8000
MAX_FILE = 5 * 1024 * 1024
IGNORED_DIRS = {".git", "__pycache__", ".venv", "venv", "node_modules"}
TEXT_EXT = (".py", ".json", ".md", ".txt", ".html", ".csv", ".js", ".ts", ".yaml", ".yml",
            ".toml", ".ini", ".cfg", ".bat", ".sh", ".c", ".cpp", ".h", ".m", ".xml")
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
      "Accept-Language": "tr-TR,tr;q=0.9,en-US;q=0.8,en;q=0.7"}


def _clip(text, limit: int = MAX_OUT) -> str:
    text = str(text)
    return text if len(text) <= limit else text[:limit] + f"\n...[{len(text) - limit} karakter kisaltildi]"


# ==========================================
# 0. GUVENLI YOL COZUMLEME (WORKSPACE KAFESI)
# ==========================================

def _safe_path(filepath: str) -> str:
    """Yolu calisma alani (workspace) icinde cozer; disina cikan her yolu reddeder.
    Goreli yol, workspace'e gore yorumlanir. '..', sembolik link ve NTFS akislari engellenir."""
    p = str(filepath or "").strip().strip("\"'")
    if not p:
        raise ValueError("Bos dosya yolu.")
    if not os.path.isabs(p):
        if ":" in p:
            raise ValueError("Yolda ':' karakterine izin verilmez.")
        parts = [x for x in re.split(r"[\\/]+", p) if x]
        if parts and parts[0].lower() == "workspace":
            parts = parts[1:]
        p = os.sep.join(parts) or "."
    full = os.path.realpath(os.path.join(WORKSPACE_DIR, p))
    if not Path(full).is_relative_to(Path(WORKSPACE_DIR)):
        raise ValueError(f"'{filepath}' calisma alani disinda; erisim engellendi.")
    return full


# ---- Docker sandbox (istege bagli: .env icinde NODUS_SANDBOX=docker) ----
SANDBOX_IMAGE = os.getenv("NODUS_SANDBOX_IMAGE", "nodus-sandbox")
_docker_state = {"t": 0.0, "ok": False, "why": ""}


def sandbox_status() -> tuple:
    """(aktif_mi, aciklama). NODUS_SANDBOX=docker, docker kurulu ve imaj hazirsa aktif."""
    if os.getenv("NODUS_SANDBOX", "").strip().lower() != "docker":
        return False, "kapali (NODUS_SANDBOX=docker ile acilir)"
    if time.time() - _docker_state["t"] < 60:
        return _docker_state["ok"], _docker_state["why"]
    ok, why = False, ""
    if not shutil.which("docker"):
        why = "Docker kurulu degil"
    else:
        try:
            r = subprocess.run(["docker", "image", "inspect", SANDBOX_IMAGE], capture_output=True, timeout=15)
            ok = r.returncode == 0
            why = "aktif (docker, ag yok)" if ok else f"'{SANDBOX_IMAGE}' imaji yok: docker build -t {SANDBOX_IMAGE} sandbox"
        except Exception as e:
            why = f"Docker yanit vermiyor ({type(e).__name__})"
    _docker_state.update(t=time.time(), ok=ok, why=why)
    return ok, why


def _docker_cmd(inner: list) -> tuple:
    """Agsiz, bellek/cpu/surec sinirli, workspace disinda hicbir seye erisemeyen konteyner komutu."""
    name = "nodus-" + uuid.uuid4().hex[:12]
    cmd = ["docker", "run", "--rm", "-i", "--name", name, "--network", "none", "--memory", "1g", "--cpus", "2",
           "--pids-limit", "256", "--cap-drop", "ALL", "--security-opt", "no-new-privileges", "--tmpfs", "/tmp",
           "-v", f"{WORKSPACE_DIR}:/work", "-w", "/work", "-e", "MPLBACKEND=Agg", "-e", "PYTHONIOENCODING=utf-8",
           SANDBOX_IMAGE] + inner
    return cmd, name


def _run(cmd, timeout: int, shell: bool = False, container: str = None):
    """Alt sureci workspace icinde, API anahtarlari olmadan calistirir. (returncode, stdout, stderr)"""
    p = subprocess.Popen(
        cmd, cwd=WORKSPACE_DIR, env=scrubbed_env(), shell=shell,
        stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        text=True, encoding="utf-8", errors="replace",
    )
    try:
        out, err = p.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        if container:  # istemciyi oldurmek konteyneri durdurmaz
            subprocess.run(["docker", "kill", container], capture_output=True)
        if os.name == "nt":  # shell=True'da cocuk surecleri de oldur
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(p.pid)], capture_output=True)
        else:
            p.kill()
        p.communicate()
        raise
    return p.returncode, out, err


def _run_py(args: list, timeout: int):
    """Python'u (sandbox aciksa konteynerde) calistirir. args: yorumlayici sonrasi arguman listesi."""
    if sandbox_status()[0]:
        cmd, name = _docker_cmd(["python"] + args)
        return _run(cmd, timeout, container=name)
    return _run([sys.executable] + args, timeout)


def _run_sh(command: str, timeout: int):
    if sandbox_status()[0]:
        cmd, name = _docker_cmd(["sh", "-c", command])
        return _run(cmd, timeout, container=name)
    return _run(command, timeout, shell=True)


def _format_run(code: int, out: str, err: str) -> str:
    if code == 0:
        return "Basarili Cikti:\n" + (_clip(out.strip()) or "Komut hatasiz calisti.") + (f"\n[stderr]\n{_clip(err.strip(), 2000)}" if err.strip() else "")
    return f"Calisma Hatasi (cikis kodu {code}):\n" + _clip((err.strip() or out.strip()))


# ==========================================
# 1. DOSYA VE DIZIN ISLEMLERI
# ==========================================

def list_files(directory: str = ".", pattern: str = "*") -> list:
    """Calisma alanindaki dosyalari (alt klasorler dahil, goreli yol) listeler."""
    try:
        root = WORKSPACE_DIR if str(directory or ".").strip() in (".", "") else _safe_path(directory)
        os.makedirs(root, exist_ok=True)
        found = []
        for dp, dn, fn in os.walk(root):
            dn[:] = [d for d in dn if d not in IGNORED_DIRS]
            for f in fn:
                rel = os.path.relpath(os.path.join(dp, f), root)
                if fnmatch.fnmatch(rel, pattern) or fnmatch.fnmatch(f, pattern):
                    found.append(rel)
                    if len(found) >= 500:
                        return sorted(found)
        return sorted(found)
    except Exception:
        return []


def read_file(filepath: str, start_line=None, end_line=None) -> str:
    """Dosyayi okur. start_line/end_line verilirse satir numarali parca doner."""
    try:
        path = _safe_path(filepath)
        if not os.path.isfile(path):
            return ERR + f"'{filepath}' dosyasi bulunamadi."
        if os.path.getsize(path) > MAX_FILE:
            return ERR + f"'{filepath}' {MAX_FILE // 1024 // 1024} MB sinirindan buyuk."
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            if not (start_line or end_line):
                return f.read()
            s = max(int(start_line or 1), 1)
            e = int(end_line) if end_line else None
            lines = []
            for i, line in enumerate(f, 1):
                if i < s:
                    continue
                if e and i > e:
                    break
                lines.append(f"{i}: {line.rstrip()}")
            return "\n".join(lines)
    except Exception as e:
        return ERR + f"Dosya okuma hatasi: {e}"


def write_file(filepath: str, content: str) -> str:
    """Calisma alaninda dosya olusturur veya uzerine yazar."""
    try:
        path = _safe_path(filepath)
        content = str(content)
        if len(content) > MAX_FILE:
            return ERR + "Icerik cok buyuk."
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return f"'{filepath}' basariyla yazildi."
    except Exception as e:
        return ERR + f"Dosya yazma hatasi: {e}"


def apply_patch(filepath: str, target_snippet: str, replacement_snippet: str) -> str:
    """Dosyada TEK bir yerde gecen blogu degistirir (birden fazla eslesme reddedilir)."""
    try:
        path = _safe_path(filepath)
        if not os.path.isfile(path):
            return ERR + f"'{filepath}' bulunamadi."
        if not target_snippet:
            return ERR + "search_block bos olamaz."

        with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
            content = f.read()

        target, repl = str(target_snippet), str(replacement_snippet)
        if target not in content and "\r\n" in content:  # Windows satir sonlari
            target, repl = target.replace("\n", "\r\n"), repl.replace("\n", "\r\n")

        count = content.count(target)
        if count == 0:
            return ERR + f"Hedef blok '{filepath}' icinde bulunamadi (bosluk/girinti birebir ayni olmali)."
        if count > 1:
            return ERR + f"Hedef blok {count} yerde bulundu; benzersiz olacak sekilde daha fazla baglam ekleyin."

        with open(path, "w", encoding="utf-8", newline="") as f:
            f.write(content.replace(target, repl, 1))
        return f"'{filepath}' basariyla yamalandi."
    except Exception as e:
        return ERR + f"Yama uygulama hatasi: {e}"


def delete_file(filepath: str) -> str:
    """Calisma alanindaki tek bir dosyayi siler."""
    try:
        path = _safe_path(filepath)
        if not os.path.isfile(path):
            return ERR + f"'{filepath}' dosya degil veya bulunamadi."
        os.remove(path)
        return f"'{filepath}' silindi."
    except Exception as e:
        return ERR + f"Silme hatasi: {e}"


def search_codebase(pattern: str, target_dir: str = ".") -> str:
    """Calisma alanindaki metin dosyalarinda buyuk/kucuk harfe duyarsiz duz metin aramasi yapar."""
    needle = str(pattern or "").strip().lower()
    if not needle:
        return ERR + "Aranacak ifade bos."
    results = []
    try:
        root = WORKSPACE_DIR if str(target_dir or ".").strip() in (".", "") else _safe_path(target_dir)
        for dp, dn, fn in os.walk(root):
            dn[:] = [d for d in dn if d not in IGNORED_DIRS]
            for file in fn:
                path = os.path.join(dp, file)
                if not file.lower().endswith(TEXT_EXT) or os.path.getsize(path) > 1024 * 1024:
                    continue
                try:
                    with open(path, "r", encoding="utf-8", errors="ignore") as f:
                        for idx, line in enumerate(f, 1):
                            if needle in line.lower():
                                results.append(f"{os.path.relpath(path, WORKSPACE_DIR)}:{idx} -> {line.strip()[:300]}")
                                if len(results) >= 50:
                                    return "\n".join(results)
                except OSError:
                    continue
        return "\n".join(results) if results else f"'{pattern}' icin eslesme bulunamadi."
    except Exception as e:
        return ERR + f"Arama hatasi: {e}"


# ==========================================
# 2. CALISTIRMA MOTORU
# ==========================================
# shortcut: alt surec yalnizca cwd/env/timeout ile sinirli, OS duzeyinde sandbox degil
# (dosya sistemini okuyabilir). Gercek yalitim icin Docker/Windows Sandbox'a tasiyin.

def execute_python_code(code: str, timeout: int = 30) -> str:
    """Python kodunu workspace icinde, API anahtarsiz alt surecte calistirir."""
    try:
        timeout = max(1, min(int(timeout), 120))
        rc, out, err = _run_py(["-X", "utf8", "-c", str(code)], timeout)
        return _format_run(rc, out, err)
    except subprocess.TimeoutExpired:
        return f"Calisma Hatasi:\nKod {timeout} saniyelik zaman asimina ugradi."
    except Exception as e:
        return f"Calisma Hatasi:\n{e}"


def run_python_file(filename: str, timeout: int = 60) -> str:
    """Workspace'teki bir .py dosyasini calistirir (onay ekraninda dosyanin icerigi gosterilir)."""
    try:
        path = _safe_path(filename)
        if not path.lower().endswith(".py") or not os.path.isfile(path):
            return ERR + f"'{filename}' mevcut bir .py dosyasi degil."
        timeout = max(1, min(int(timeout), 300))
        rc, out, err = _run_py(["-X", "utf8", os.path.relpath(path, WORKSPACE_DIR)], timeout)
        return _format_run(rc, out, err)
    except subprocess.TimeoutExpired:
        return f"Calisma Hatasi:{chr(10)}Dosya {timeout} saniyelik zaman asimina ugradi."
    except ValueError as e:  # yol kafesi ihlali
        return ERR + str(e)
    except Exception as e:
        return f"Calisma Hatasi:{chr(10)}{e}"


def run_shell(command: str, timeout: int = 60) -> str:
    """Kabuk komutu calistirir (workspace icinde). Her zaman kullanici onayi gerektirir."""
    try:
        timeout = max(1, min(int(timeout), 300))
        rc, out, err = _run_sh(str(command), timeout)
        return _format_run(rc, out, err)
    except subprocess.TimeoutExpired:
        return f"Calisma Hatasi:\nKomut {timeout} saniyelik zaman asimina ugradi."
    except Exception as e:
        return f"Calisma Hatasi:\n{e}"


# ==========================================
# 3. INTERNET: ARAMA, SAYFA OKUMA, EKRAN
# ==========================================

def web_search(query: str, max_results: int = 5) -> str:
    """Canli internet aramasi (ddgs; yedek: googlesearch). Sonuclarda URL de doner."""
    output = []
    try:  # ddgs, Bing/Brave/DuckDuckGo arka uclarini dener (DDG Lite artik bot engeli veriyor)
        from ddgs import DDGS
        for r in DDGS().text(query, max_results=max_results):
            output.append(f"- {r.get('title', '')}: {(r.get('body') or '').strip()} ({r.get('href', '')})")
    except Exception:
        pass

    if not output:
        try:
            from googlesearch import search
            for r in search(query, num_results=max_results, advanced=True, lang="tr"):
                output.append(f"- {r.title}: {r.description} ({r.url})")
        except Exception:
            pass

    if not output:
        return f"'{query}' aramasi icin internette sonuc bulunamadi."
    return "\n".join(output[:max_results])


def _assert_public_url(url: str):
    u = urlparse(url)
    if u.scheme not in ("http", "https") or not u.hostname:
        raise ValueError("Yalnizca http/https URL'leri desteklenir.")
    port = u.port or (443 if u.scheme == "https" else 80)
    for info in socket.getaddrinfo(u.hostname, port):
        ip = ipaddress.ip_address(info[4][0].split("%")[0])
        if ip.version == 6 and ip.ipv4_mapped:
            ip = ip.ipv4_mapped
        if not ip.is_global:
            raise ValueError("Yerel/ozel ag adreslerine erisim engellendi (SSRF korumasi).")


def fetch_url(url: str, max_chars: int = MAX_OUT) -> str:
    """Bir web sayfasini indirip okunabilir metne cevirir. Yerel/ozel adresler engellenir."""
    # shortcut: DNS kontrolu ile istek arasinda kucuk bir rebinding penceresi var; proxy/egress filtresi ekleyin.
    try:
        url = str(url).strip()
        for _ in range(4):
            _assert_public_url(url)
            r = requests.get(url, headers={**UA, "User-Agent": "NodusBot/1.0 (yerel arastirma asistani)"}, timeout=(5, 15), stream=True, allow_redirects=False)
            if r.status_code in (301, 302, 303, 307, 308):
                url = urljoin(url, r.headers.get("Location", ""))
                r.close()
                continue
            break
        else:
            return ERR + "Cok fazla yonlendirme."

        if r.status_code >= 400:
            r.close()
            return ERR + f"Sunucu {r.status_code} dondurdu."
        ctype = r.headers.get("Content-Type", "").lower()
        if ctype and not any(t in ctype for t in ("text", "json", "xml")):
            r.close()
            return ERR + f"Desteklenmeyen icerik turu: {ctype}"

        buf = b""
        for chunk in r.iter_content(65536):
            buf += chunk
            if len(buf) > 1_000_000:
                break
        r.close()

        text = buf.decode(r.encoding or "utf-8", errors="replace")
        if "html" in ctype or text.lstrip()[:15].lower().startswith(("<!doctype", "<html")):
            text = re.sub(r"(?is)<(script|style|noscript|svg)\b.*?</\1>", " ", text)
            text = unescape(re.sub(r"<[^>]+>", " ", text))
        text = re.sub(r"[ \t\r\f\v]+", " ", text)
        text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
        return _clip(text, max(500, min(int(max_chars), 30000))) or "(Sayfa bos.)"
    except Exception as e:
        return ERR + f"Sayfa alinamadi: {e}"


def capture_screen(output_path: str = "screen_capture.png") -> str:
    """Masaustunun anlik ekran goruntusunu workspace'e kaydeder."""
    try:
        import mss
        path = _safe_path(output_path)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with mss.mss() as sct:
            sct.shot(mon=-1, output=path)
        return f"Ekran goruntusu kaydedildi: {output_path}"
    except Exception as e:
        return ERR + f"Ekran yakalama hatasi: {e}"


# ==========================================
# 4. SEMBOLIK MATEMATIK (EVAL-GUVENLI)
# ==========================================

_SAFE_EXPR = re.compile(r"^[\w\s+\-*/^().,=\[\]]+$")
_SYMPY_TASKS = {"solve", "diff", "integrate", "simplify", "expand", "factor"}

# Alt surecte calisir: eval'e acik __builtins__ verilmez; ifade argv ile gelir.
_SYMPY_CODE = r'''
import sys
import sympy as sp
from sympy.parsing.sympy_parser import parse_expr, standard_transformations, convert_xor, implicit_multiplication_application
expr_s, var_s, task = sys.argv[1:4]
g = {}
exec("from sympy import *", g)
g["__builtins__"] = {}
tr = standard_transformations + (convert_xor, implicit_multiplication_application)
P = lambda s: parse_expr(s, global_dict=g, transformations=tr)
var = sp.Symbol(var_s)
if "=" in expr_s:
    l, r = expr_s.split("=", 1)
    expr = P(l) - P(r)
else:
    expr = P(expr_s)
if task == "solve":
    print("Denklem: %s = 0" % expr)
    print("Cozum kumesi: %s" % (sp.solve(expr, var),))
elif task in ("diff", "integrate"):
    print("Sonuc (%s): %s" % (task, getattr(sp, task)(expr, var)))
else:
    print("Sonuc (%s): %s" % (task, getattr(sp, task)(expr)))
'''


def solve_symbolic_math(equation_str: str, variable_str: str = "x", task: str = "solve") -> str:
    """SymPy ile denklem cozer / turev / integral / sadelestirme yapar (izole alt surecte)."""
    try:
        expr = str(equation_str or "").strip().replace("==", "=")
        var = str(variable_str or "x").strip()
        task = str(task or "solve").strip().lower()
        if (not expr or len(expr) > 500 or not _SAFE_EXPR.match(expr)
                or "__" in expr or re.search(r"\.\s*[A-Za-z_]", expr)):
            return ERR + "Gecersiz veya izin verilmeyen ifade (nokta ile nitelik, tirnak, ':' vb. yasak)."
        if not re.fullmatch(r"[A-Za-z_]\w{0,15}", var):
            return ERR + "Gecersiz degisken adi."
        if task not in _SYMPY_TASKS:
            return ERR + f"Bilinmeyen gorev '{task}'. Gecerli: {sorted(_SYMPY_TASKS)}"
        rc, out, err = _run([sys.executable, "-X", "utf8", "-c", _SYMPY_CODE, expr, var, task], 30)
        if rc != 0:
            return ERR + "Sembolik cozum hatasi: " + _clip(err.strip().splitlines()[-1] if err.strip() else "bilinmiyor", 500)
        return out.strip()
    except subprocess.TimeoutExpired:
        return ERR + "Sembolik cozum zaman asimina ugradi."
    except Exception as e:
        return ERR + f"Sembolik cozum hatasi: {e}"


# ==========================================
# 5. OTONOM MODUL GELISTIRICI (self_improve)
# ==========================================

def _strip_fences(code: str) -> str:
    code = code.strip()
    code = re.sub(r"^```[a-zA-Z]*\s*\n", "", code)
    return re.sub(r"\n```\s*$", "", code).strip()


def _ask_llm(prompt: str) -> str:
    """Once yerel Ollama, olmazsa OpenRouter ucretsiz modeli."""
    try:
        r = requests.post("http://localhost:11434/api/generate",
                          json={"model": "minicpm-v", "prompt": prompt, "stream": False}, timeout=120)
        if r.status_code == 200 and r.json().get("response"):
            return r.json()["response"]
    except Exception:
        pass
    key = os.getenv("OPENROUTER_API_KEY")
    if not key:
        return ""
    for model in ("meta-llama/llama-3.3-70b-instruct:free", "deepseek/deepseek-r1:free"):
        try:
            r = requests.post("https://openrouter.ai/api/v1/chat/completions",
                              headers={"Authorization": f"Bearer {key}"},
                              json={"model": model, "messages": [{"role": "user", "content": prompt}], "temperature": 0.2},
                              timeout=(10, 120))
            if r.status_code == 200:
                return r.json()["choices"][0]["message"]["content"] or ""
        except Exception:
            continue
    return ""


def self_improve(task_description: str) -> str:
    """Yeni araci custom_tools/<kategori>.py dosyasina yazar, izole test eder ve yukler.
    tools.py'ye dokunmaz. Kullanici onayi gerektirir (kod bu surece yuklenir)."""
    test_script = os.path.join(BASE_DIR, "test_candidate.py")

    prompt = (
        "Sen otonom bir makine muhendisligi yazilim mimarisisin.\n\n"
        "GOREV: " + str(task_description) + "\n\n"
        "Bu gorev icin bir Python modulu yazacaksin. tools.py dosyasina dokunulmayacak.\n"
        "Gorevin ait oldugu kategoriyi belirle (ornek: fluids, mechanics, materials, thermal, math).\n"
        "Fonksiyonlar tipli parametre alsin, docstring icersin, yalnizca standart kutuphane/math kullansin.\n\n"
        "Lutfen su 3 etiketi eksiksiz uret:\n"
        "<CATEGORY>\nkategori_adi (sadece tek kelime ingilizce harfler, ornegin: fluids)\n</CATEGORY>\n"
        "<CODE>\n# Kategori dosyasina eklenecek tam fonksiyon kodu\n</CODE>\n"
        "<TEST>\n# Bu fonksiyonu 'from custom_tools.<kategori> import ...' ile test eden kod. sys.exit(0) ile bitmeli\n</TEST>"
    )

    response_text = _ask_llm(prompt)
    if not response_text:
        return ERR + "Model erisim hatasi (Ollama ve OpenRouter yanit vermedi)."

    cat_match = re.search(r"<CATEGORY>(.*?)</CATEGORY>", response_text, re.DOTALL)
    code_match = re.search(r"<CODE>(.*?)</CODE>", response_text, re.DOTALL)
    test_match = re.search(r"<TEST>(.*?)</TEST>", response_text, re.DOTALL)
    if not (cat_match and code_match and test_match):
        return ERR + "Model istenen formatta kod/test uretemedi. Islem durduruldu."

    category = re.sub(r"[^a-z0-9_]", "", cat_match.group(1).strip().lower().replace(".py", "")) or "general"
    new_code = _strip_fences(code_match.group(1))
    test_code = _strip_fences(test_match.group(1))

    target_file = os.path.join(CUSTOM_DIR, f"{category}.py")
    old_content = None
    if os.path.exists(target_file):
        with open(target_file, "r", encoding="utf-8") as f:
            old_content = f.read()
        final_code = old_content + "\n\n" + new_code
    else:
        final_code = "# -*- coding: utf-8 -*-\n" + new_code

    with open(target_file, "w", encoding="utf-8") as f:
        f.write(final_code)
    with open(test_script, "w", encoding="utf-8") as f:
        f.write(test_code)

    try:
        success = BackupGuard.run_candidate_test(test_script, target_file)
    finally:
        if os.path.exists(test_script):
            os.remove(test_script)

    if success:
        load_dynamic_tools()
        return f"Basarili! Yeni fonksiyon '{target_file}' dosyasina eklendi ve yuklendi.\n\nEklenen kod:\n```python\n{new_code}\n```"

    if old_content is not None:
        with open(target_file, "w", encoding="utf-8") as f:
            f.write(old_content)
    elif os.path.exists(target_file):
        os.remove(target_file)
    return ERR + f"Test basarisiz oldu. '{target_file}' dosyasina yapilan ekleme iptal edildi."


# ==========================================
# 5a. BELGE OKUMA, RAPOR URETIMI, TARAYICI
# ==========================================

MAX_DOC = 20 * 1024 * 1024
DOC_EXT = (".csv", ".txt", ".md", ".json", ".xlsx", ".pdf", ".docx")


def read_file_bytes(filepath: str) -> bytes:
    """Ikili dosyayi (indirme icin) workspace icinden okur; hata durumunda b''."""
    try:
        path = _safe_path(filepath)
        if os.path.isfile(path) and os.path.getsize(path) <= MAX_DOC:
            with open(path, "rb") as f:
                return f.read()
    except Exception:
        pass
    return b""


def save_upload(filename: str, data: bytes) -> tuple:
    """Kullanicinin yukledigi belgeyi workspace/uploads altina kaydeder. (ok, 'uploads/ad.ext' | hata)."""
    try:
        name = re.sub(r"[^\w.\-]", "_", os.path.basename(str(filename or "")))[:80].lstrip(".")
        ext = os.path.splitext(name)[1].lower()
        if ext not in DOC_EXT:
            return False, f"Desteklenmeyen dosya turu '{ext}'. Desteklenen: {', '.join(DOC_EXT)}"
        if len(data) > MAX_DOC:
            return False, "Dosya 20 MB sinirindan buyuk."
        stem, n, rel = name[:-len(ext)], 1, f"uploads/{name}"
        while os.path.exists(_safe_path(rel)):      # var olan dosyanin uzerine yazma
            n += 1
            rel = f"uploads/{stem}_{n}{ext}"
        path = _safe_path(rel)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "wb") as f:
            f.write(data)
        return True, rel
    except Exception as e:
        return False, f"Dosya kaydedilemedi: {e}"


def read_document(filename: str, max_chars: int = 12000) -> str:
    """CSV/TXT/JSON/XLSX/PDF/DOCX belgelerinin metnini okur (yuklenen dosyalar workspace/uploads altindadir)."""
    try:
        path = _safe_path(filename)
        ext = os.path.splitext(path)[1].lower()
        if not os.path.isfile(path):
            return ERR + f"'{filename}' bulunamadi."
        if os.path.getsize(path) > MAX_DOC:
            return ERR + "Belge 20 MB sinirindan buyuk."
        if ext in (".csv", ".txt", ".md", ".json"):
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read(MAX_DOC)
        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets[:5]:
                parts.append(f"## Sayfa: {ws.title}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 60:
                        parts.append("...(satirlar kisaltildi)")
                        break
                    parts.append(" | ".join("" if c is None else str(c) for c in row))
            wb.close()
            text = "\n".join(parts)
        elif ext == ".pdf":
            from pypdf import PdfReader
            rd = PdfReader(path)
            if rd.is_encrypted:
                return ERR + "Sifreli PDF desteklenmiyor."
            text = "\n".join(f"[Sayfa {i}]\n{(pg.extract_text() or '').strip()}" for i, pg in enumerate(rd.pages[:30], 1))
        elif ext == ".docx":
            import docx
            d = docx.Document(path)
            text = "\n".join(p.text for p in d.paragraphs)
            for t in d.tables:
                text += "\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows)
        else:
            return ERR + f"Desteklenmeyen belge turu '{ext}'. Desteklenen: {', '.join(DOC_EXT)}"
        return _clip(text.strip() or "(belge bos veya metin cikarilamadi)", max(500, min(int(max_chars), 30000)))
    except Exception as e:
        return ERR + f"Belge okunamadi: {type(e).__name__}: {e}"


def _parse_md(content: str) -> list:
    """Markdown alt kumesi -> bloklar: h, p, ul, ol, table, img."""
    blocks, table = [], []

    def flush():
        nonlocal table
        if table:
            blocks.append(("table", table))
            table = []

    for raw in str(content).splitlines():
        line = raw.rstrip()
        if re.match(r"^\s*\|.*\|\s*$", line):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):   # ayirici satiri atla
                table.append(cells)
            continue
        flush()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            blocks.append(("h", len(m.group(1)), m.group(2)))
            continue
        m = re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", line)
        if m:
            blocks.append(("img", m.group(2), m.group(1)))
            continue
        m = re.match(r"^\s*[-*]\s+(.*)", line)
        if m:
            blocks.append(("ul", m.group(1)))
            continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if m:
            blocks.append(("ol", m.group(1)))
            continue
        blocks.append(("p", line.strip()))
    flush()
    return blocks


def _report_image(ref: str):
    """Rapora gomulecek gorsel: yalnizca workspace icinde, png/jpg."""
    try:
        p = _safe_path(ref)
        return p if os.path.isfile(p) and p.lower().endswith((".png", ".jpg", ".jpeg")) else None
    except ValueError:
        return None


def _render_docx(title, blocks, out):
    import docx
    from docx.shared import Inches
    d = docx.Document()
    d.add_heading(title, 0)

    def runs(par, text):
        for i, part in enumerate(text.split("**")):
            par.add_run(part).bold = (i % 2 == 1)

    for b in blocks:
        if b[0] == "h":
            d.add_heading(b[2].replace("**", ""), min(b[1], 3))
        elif b[0] == "p":
            runs(d.add_paragraph(), b[1])
        elif b[0] == "ul":
            runs(d.add_paragraph(style="List Bullet"), b[1])
        elif b[0] == "ol":
            runs(d.add_paragraph(style="List Number"), b[1])
        elif b[0] == "table":
            t = d.add_table(rows=len(b[1]), cols=max(len(r) for r in b[1]))
            t.style = "Table Grid"
            for i, row in enumerate(b[1]):
                for j, c in enumerate(row):
                    t.cell(i, j).text = c.replace("**", "")
        elif b[0] == "img":
            path = _report_image(b[1])
            if path:
                d.add_picture(path, width=Inches(6))
            else:
                runs(d.add_paragraph(), f"[gorsel bulunamadi: {b[1]}]")
    d.save(out)


def _render_pdf(title, blocks, out):
    import matplotlib
    from fpdf import FPDF
    fonts = os.path.join(os.path.dirname(matplotlib.__file__), "mpl-data", "fonts", "ttf")   # Turkce karakterli TTF
    pdf = FPDF()
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    pdf.add_font("DejaVu", "", os.path.join(fonts, "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", os.path.join(fonts, "DejaVuSans-Bold.ttf"))

    def write(text, size=11, bold=False, indent=0):
        pdf.set_font("DejaVu", "B" if bold else "", size)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(0, size * 0.55, text.replace("**", ""), new_x="LMARGIN", new_y="NEXT")

    write(title, 18, True)
    pdf.ln(2)
    n = 0
    for b in blocks:
        if b[0] != "ol":
            n = 0
        if b[0] == "h":
            pdf.ln(2)
            write(b[2], {1: 15, 2: 13}.get(b[1], 12), True)
        elif b[0] == "p":
            write(b[1])
        elif b[0] == "ul":
            write("• " + b[1], indent=4)
        elif b[0] == "ol":
            n += 1
            write(f"{n}. " + b[1], indent=4)
        elif b[0] == "table":
            pdf.set_font("DejaVu", "", 10)
            with pdf.table() as table:
                for row in b[1]:
                    r = table.row()
                    for c in row:
                        r.cell(c.replace("**", ""))
            pdf.ln(2)
        elif b[0] == "img":
            path = _report_image(b[1])
            if path:
                pdf.image(path, w=170)
                pdf.ln(3)
            else:
                write(f"[gorsel bulunamadi: {b[1]}]")
    pdf.output(out)


def _render_html(title, blocks, out):
    from html import escape
    body = [f"<h1>{escape(title)}</h1>"]
    for b in blocks:
        if b[0] == "h":
            body.append(f"<h{b[1] + 1}>{escape(b[2])}</h{b[1] + 1}>")
        elif b[0] in ("p", "ul", "ol"):
            txt = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escape(b[1]))
            body.append(f"<p>{txt}</p>" if b[0] == "p" else f"<p>&bull; {txt}</p>")
        elif b[0] == "table":
            rows = "".join("<tr>" + "".join(f"<td>{escape(c)}</td>" for c in r) + "</tr>" for r in b[1])
            body.append(f"<table border='1' cellpadding='4' style='border-collapse:collapse'>{rows}</table>")
        elif b[0] == "img":
            path = _report_image(b[1])
            if path:
                import base64
                with open(path, "rb") as f:
                    mime = "image/png" if path.lower().endswith(".png") else "image/jpeg"
                    body.append(f"<img style='max-width:100%' src='data:{mime};base64,{base64.b64encode(f.read()).decode()}'>")
    with open(out, "w", encoding="utf-8") as f:
        f.write("<!doctype html><meta charset='utf-8'><body style='font-family:Arial;max-width:800px;margin:2em auto'>" + "".join(body))


def export_report(filename: str, title: str, content: str, format: str = "docx") -> str:
    """Markdown icerikten rapor uretir: docx | pdf | html | md.  Gorsel: ![aciklama](grafik.png)  (workspace icinden)."""
    try:
        fmt = str(format or "docx").lower().strip(".")
        if fmt not in ("docx", "pdf", "html", "md"):
            return ERR + "format docx, pdf, html veya md olmali."
        content = str(content or "")
        if not content.strip() or len(content) > 200_000:
            return ERR + "Rapor icerigi bos veya cok buyuk."
        base = re.sub(r"[^\w\-]", "_", os.path.splitext(os.path.basename(str(filename or "rapor")))[0]) or "rapor"
        out = _safe_path(f"reports/{base}.{fmt}")
        os.makedirs(os.path.dirname(out), exist_ok=True)
        title = str(title or "Nodus Raporu")[:200]
        blocks = _parse_md(content)
        if fmt == "docx":
            _render_docx(title, blocks, out)
        elif fmt == "pdf":
            _render_pdf(title, blocks, out)
        elif fmt == "html":
            _render_html(title, blocks, out)
        else:
            with open(out, "w", encoding="utf-8") as f:
                f.write(f"# {title}\n\n{content}")
        return f"Rapor kaydedildi: reports/{base}.{fmt}"
    except Exception as e:
        return ERR + f"Rapor olusturulamadi: {type(e).__name__}: {e}"


def browse_url(url: str, mode: str = "text") -> str:
    """Gercek tarayicida (JavaScript calistirarak) sayfayi acar. mode='text' | 'screenshot'.
    Ozel/yerel adresler her istekte engellenir. fetch_url yetersiz kalan dinamik sayfalar icindir."""
    try:
        url = str(url).strip()
        _assert_public_url(url)
        shot = ""
        if str(mode).lower().startswith("s"):
            shot = _safe_path(f"browser/{uuid.uuid4().hex[:10]}.png")
            os.makedirs(os.path.dirname(shot), exist_ok=True)
        rc, out, err = _run([sys.executable, "-X", "utf8", os.path.join(BASE_DIR, "browse_worker.py"), url,
                             "screenshot" if shot else "text", shot], 60)
        if rc != 0:
            last = err.strip().splitlines()[-1] if err.strip() else "bilinmiyor"
            return ERR + "Tarayici hatasi: " + _clip(last, 300)
        if shot:
            return "Ekran goruntusu kaydedildi: " + os.path.relpath(shot, WORKSPACE_DIR)
        return _clip(out.strip() or "(sayfa bos)")
    except subprocess.TimeoutExpired:
        return ERR + "Tarayici 60 saniyede yanit vermedi."
    except Exception as e:
        return ERR + f"Sayfa acilamadi: {e}"


# ==========================================
# 5b. PLAN VE KALICI NOT HAFIZASI
# ==========================================

PLAN = []            # [{"text": ..., "status": pending|in_progress|done}]
_PLAN_ICONS = {"pending": "☐", "in_progress": "◐", "done": "☑"}
_notes = None


def format_plan() -> str:
    return "\n".join(f"{_PLAN_ICONS[p['status']]} {i}. {p['text']}" for i, p in enumerate(PLAN, 1)) or "(plan bos)"


def update_plan(steps) -> str:
    """Cok adimli gorevin planini gunceller. steps: [{"text": "...", "status": "pending|in_progress|done"}, ...]"""
    try:
        if isinstance(steps, str):
            try:
                steps = json.loads(steps)
            except ValueError:
                steps = [x for x in steps.splitlines() if x.strip()]
        if not isinstance(steps, list) or not steps:
            return ERR + "steps bos olamaz; adim listesi verin."
        new = []
        for st in steps[:15]:
            text = str(st.get("text") or st.get("step") or "") if isinstance(st, dict) else str(st)
            status = str(st.get("status", "pending")).lower() if isinstance(st, dict) else "pending"
            if text.strip():
                new.append({"text": text.strip()[:200], "status": status if status in _PLAN_ICONS else "pending"})
        PLAN[:] = new
        return "Plan guncellendi:\n" + format_plan()
    except Exception as e:
        return ERR + f"Plan hatasi: {e}"


def _get_notes():
    global _notes
    if _notes is None:
        from memory_manager import NodusNotes
        _notes = NodusNotes()
    return _notes


def remember(text: str, tags: str = "") -> str:
    """Kalici uzun sureli nota yazar (oturumlar arasi hatirlanir)."""
    try:
        return "Not kaydedildi." if _get_notes().add(text, tags) else ERR + "Bos not kaydedilmedi."
    except Exception as e:
        return ERR + f"Not kaydi hatasi: {e}"


def recall(query: str, limit: int = 5) -> str:
    """Kalici notlarda tam metin arar."""
    try:
        hits = _get_notes().search(query, max(1, min(int(limit), 10)))
        if not hits:
            return f"'{query}' icin kayitli not bulunamadi."
        return "\n".join(f"- [{h['created']}] {h['text']}" + (f" (etiket: {h['tags']})" if h["tags"] else "") for h in hits)
    except Exception as e:
        return ERR + f"Not arama hatasi: {e}"


# ==========================================
# 6. ARAC KAYIT DEFTERI
# ==========================================

TOOL_REGISTRY = {
    "list_files": list_files,
    "read_file": read_file,
    "write_file": write_file,
    "apply_patch": apply_patch,
    "delete_file": delete_file,
    "search_codebase": search_codebase,
    "execute_python_code": execute_python_code,
    "run_shell": run_shell,
    "run_python_file": run_python_file,
    "web_search": web_search,
    "fetch_url": fetch_url,
    "capture_screen": capture_screen,
    "solve_symbolic_math": solve_symbolic_math,
    "self_improve": self_improve,
    "read_document": read_document,
    "export_report": export_report,
    "browse_url": browse_url,
    "update_plan": update_plan,
    "remember": remember,
    "recall": recall,
}
CORE_TOOLS = frozenset(TOOL_REGISTRY)
CUSTOM_TOOLS = {}


def load_dynamic_tools():
    """custom_tools/*.py icindeki YALNIZCA o modulde tanimli genel fonksiyonlari yukler.
    Cekirdek araclarin uzerine yazilamaz; ice aktarilan (os.system gibi) cagrilabilirler alinmaz."""
    for name in list(CUSTOM_TOOLS):
        TOOL_REGISTRY.pop(name, None)
    CUSTOM_TOOLS.clear()

    for fname in sorted(os.listdir(CUSTOM_DIR)):
        if not fname.endswith(".py") or fname.startswith("__"):
            continue
        mod_name = fname[:-3]
        try:
            spec = importlib.util.spec_from_file_location(mod_name, os.path.join(CUSTOM_DIR, fname))
            mod = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(mod)
            for attr_name, attr in vars(mod).items():
                if (not attr_name.startswith("_") and inspect.isfunction(attr)
                        and attr.__module__ == mod_name and attr_name not in CORE_TOOLS):
                    CUSTOM_TOOLS[attr_name] = attr
        except Exception as e:
            print(f"[LOADER] {fname} yuklenirken hata: {e}")
    TOOL_REGISTRY.update(CUSTOM_TOOLS)


load_dynamic_tools()


def custom_tool_summaries() -> list:
    """Sistem promptu icin: ['ad(imza) - ilk docstring satiri', ...]"""
    out = []
    for name, fn in CUSTOM_TOOLS.items():
        doc = (inspect.getdoc(fn) or "").strip().splitlines()
        out.append(f"{name}{inspect.signature(fn)} - {doc[0] if doc else 'aciklama yok'}")
    return out


def call_custom_tool(name: str, args: dict = None) -> str:
    """self_improve ile eklenmis ozel bir araci adiyla cagirir."""
    fn = CUSTOM_TOOLS.get(str(name))
    if not fn:
        return ERR + f"'{name}' adinda ozel arac yok. Mevcut: {list(CUSTOM_TOOLS) or 'hicbiri'}"
    try:
        return _clip(fn(**(args if isinstance(args, dict) else {})))
    except Exception as e:
        return ERR + f"Ozel arac hatasi ({name}): {e}"


def call_tool(name: str, **kwargs) -> str:
    """Ajanin metin tabanli arac cagrilari icin kopru (dahili kullanim)."""
    if name not in TOOL_REGISTRY:
        return ERR + f"Tanimlanmamis arac '{name}'."
    try:
        return TOOL_REGISTRY[name](**kwargs)
    except Exception as e:
        return ERR + f"Arac calistirma hatasi ({name}): {e}"


class NodusTools:
    """agent_core ve app tarafindan dogrudan cagirilan arac sinifi."""
    workspace_dir = WORKSPACE_DIR
    list_files = staticmethod(list_files)
    read_file = staticmethod(read_file)
    write_file = staticmethod(write_file)
    apply_patch = staticmethod(apply_patch)
    delete_file = staticmethod(delete_file)
    search_codebase = staticmethod(search_codebase)
    execute_python_code = staticmethod(execute_python_code)
    execute_code = staticmethod(execute_python_code)
    run_code = staticmethod(execute_python_code)
    run_shell = staticmethod(run_shell)
    run_python_file = staticmethod(run_python_file)
    web_search = staticmethod(web_search)
    fetch_url = staticmethod(fetch_url)
    capture_screen = staticmethod(capture_screen)
    solve_symbolic_math = staticmethod(solve_symbolic_math)
    self_improve = staticmethod(self_improve)
    call_custom_tool = staticmethod(call_custom_tool)
    read_document = staticmethod(read_document)
    save_upload = staticmethod(save_upload)
    read_file_bytes = staticmethod(read_file_bytes)
    export_report = staticmethod(export_report)
    browse_url = staticmethod(browse_url)
    sandbox_status = staticmethod(sandbox_status)
    update_plan = staticmethod(update_plan)
    remember = staticmethod(remember)
    recall = staticmethod(recall)
    format_plan = staticmethod(format_plan)
    plan = PLAN
    custom_tool_summaries = staticmethod(custom_tool_summaries)
    call_tool = staticmethod(call_tool)
    registry = TOOL_REGISTRY
